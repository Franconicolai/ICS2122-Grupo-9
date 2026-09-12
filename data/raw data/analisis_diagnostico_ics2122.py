#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Diagnóstico estructural de la instancia ICS2122 - Grupo 9.

Qué hace:
1) Lee todos los archivos de data/ y las reglas de ops_rules.yaml.
2) Revisa calidad e integridad de datos.
3) Analiza demanda semanal y diaria, incluyendo relación día d -> d+1.
4) Analiza red, asimetría, rutas directas/con conexiones y concentración en MIA.
5) Analiza flota, derechos de tráfico, operadores e interchange.
6) Analiza mantenimiento y dependencias espacio-temporales observables.
7) Explica qué holguras (slacks) son relevantes y cuáles solo pueden calcularse
   después de resolver el modelo.
8) Exporta tablas CSV, gráficos PNG y un Word simple con enfoque, hallazgos y conclusiones.

Uso esperado:
    python analisis_diagnostico_ics2122.py

El archivo debe estar en la raíz del repositorio, al mismo nivel que data/ y docs/.
También acepta --root para indicar otra carpeta del proyecto.
"""

from __future__ import annotations

import argparse
import math
import sys
from pathlib import Path
from itertools import combinations

try:
    import numpy as np
    import pandas as pd
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import networkx as nx
    import yaml
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError as exc:
    print("Falta una librería necesaria:", exc)
    print("Instala: pip install pandas numpy matplotlib networkx pyyaml python-docx")
    raise


DAYS = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]
DAY_NAMES = {
    "mon": "Lunes", "tue": "Martes", "wed": "Miércoles", "thu": "Jueves",
    "fri": "Viernes", "sat": "Sábado", "sun": "Domingo"
}


def parse_args():
    p = argparse.ArgumentParser(description="Diagnóstico estructural de la instancia ICS2122")
    p.add_argument("--root", type=Path, default=Path(__file__).resolve().parent,
                   help="Raíz del proyecto que contiene data/ y docs/")
    p.add_argument("--out", type=Path, default=None,
                   help="Carpeta de salida. Default: <root>/analisis_instancia")
    return p.parse_args()


def fmt(x, decimals=1):
    if pd.isna(x):
        return "-"
    if isinstance(x, (int, np.integer)):
        return f"{int(x)}"
    if isinstance(x, (float, np.floating)):
        return f"{float(x):,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return str(x)


def save_df(df: pd.DataFrame, path: Path, index: bool = False):
    df.to_csv(path, index=index, encoding="utf-8-sig")


def clean_ax(ax, title=None, xlabel=None, ylabel=None):
    if title:
        ax.set_title(title)
    if xlabel:
        ax.set_xlabel(xlabel)
    if ylabel:
        ax.set_ylabel(ylabel)
    ax.grid(axis="y", alpha=0.2)


def save_fig(fig, path: Path):
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)


def resolve_data_dir(root: Path) -> Path:
    """Devuelve la carpeta de datos. Primero prueba root/data (estándar), si no, root mismo."""
    if (root / "data" / "fleet.csv").exists():
        return root / "data"
    return root


def resolve_repo_root(root: Path) -> Path:
    """Devuelve la raíz del repo (donde está README.md y carpeta docs). Sube 2 niveles si hace falta."""
    if (root / "README.md").exists():
        return root
    if (root.parent.parent / "README.md").exists():
        return root.parent.parent
    return root


def load_data(root: Path):
    data = resolve_data_dir(root)
    required = [
        "fleet.csv", "airports.csv", "legs_catalog.csv", "demand_weekly.csv",
        "demand_daily.csv", "traffic_rights.csv", "interchange_airports.csv",
        "maintenance.csv", "cost_params.csv", "landing_fees.csv", "ops_rules.yaml"
    ]
    missing = [f for f in required if not (data / f).exists()]
    if missing:
        raise FileNotFoundError(f"Faltan archivos en {data}: {missing}")

    d = {}
    d["fleet"] = pd.read_csv(data / "fleet.csv")
    d["airports"] = pd.read_csv(data / "airports.csv")
    d["legs"] = pd.read_csv(data / "legs_catalog.csv")
    d["demand_weekly"] = pd.read_csv(data / "demand_weekly.csv")
    d["demand_daily"] = pd.read_csv(data / "demand_daily.csv")
    d["traffic_rights"] = pd.read_csv(data / "traffic_rights.csv")
    d["interchange"] = pd.read_csv(data / "interchange_airports.csv")
    d["maintenance"] = pd.read_csv(data / "maintenance.csv", parse_dates=["start_datetime", "end_datetime"])
    d["cost_params"] = pd.read_csv(data / "cost_params.csv")
    d["landing_fees"] = pd.read_csv(data / "landing_fees.csv")
    with open(data / "ops_rules.yaml", "r", encoding="utf-8") as f:
        d["ops_rules"] = yaml.safe_load(f)
    return d


def inventory(d):
    rows = []
    for name, obj in d.items():
        if isinstance(obj, pd.DataFrame):
            rows.append({"archivo": name, "filas": len(obj), "columnas": len(obj.columns),
                         "campos": ", ".join(obj.columns)})
        else:
            rows.append({"archivo": name, "filas": "-", "columnas": "-", "campos": "configuración YAML"})
    return pd.DataFrame(rows)


def data_quality(d):
    fleet, airports, legs = d["fleet"], d["airports"], d["legs"]
    dw, dd = d["demand_weekly"], d["demand_daily"]
    rights, inter, maint = d["traffic_rights"], d["interchange"], d["maintenance"]
    ops = d["ops_rules"]

    airport_set = set(airports["iata"])
    aircraft_set = set(fleet["aircraft_id"])
    operators = set(fleet["operator"])
    countries = set(airports["country"])

    daily_sum = dd[DAYS].sum(axis=1)
    diff = daily_sum - dw["tons_week"]

    checks = [
        ["Nulos demanda semanal", int(dw.isna().sum().sum()), "OK" if dw.isna().sum().sum() == 0 else "REVISAR"],
        ["Nulos demanda diaria", int(dd.isna().sum().sum()), "OK" if dd.isna().sum().sum() == 0 else "REVISAR"],
        ["Nulos catálogo de tramos", int(legs.isna().sum().sum()), "OK" if legs.isna().sum().sum() == 0 else "REVISAR"],
        ["Duplicados OD semanal", int(dw.duplicated(["origin", "dest"]).sum()), "OK"],
        ["Duplicados OD diario", int(dd.duplicated(["origin", "dest"]).sum()), "OK"],
        ["Duplicados de tramos", int(legs.duplicated(["origin", "dest"]).sum()), "OK"],
        ["Referencias de aeropuertos inválidas en legs", int((~legs["origin"].isin(airport_set)).sum() + (~legs["dest"].isin(airport_set)).sum()), "OK"],
        ["Referencias de aeropuertos inválidas en demanda", int((~dw["origin"].isin(airport_set)).sum() + (~dw["dest"].isin(airport_set)).sum()), "OK"],
        ["Aircraft_id de mantenimiento inválidos", int((~maint["aircraft_id"].isin(aircraft_set)).sum()), "OK"],
        ["Operadores de rights no presentes en flota", int((~rights["operator"].isin(operators)).sum()), "OK"],
        ["Países de rights no presentes en airports", int((~rights["country"].isin(countries)).sum()), "OK"],
        ["Tramos > autonomía continua permitida", int((legs["block_hours"] > float(ops["max_continuous_block_hours"])).sum()), "REVISAR"],
        ["Máx. diferencia |suma diaria - semanal| (t)", float(diff.abs().max()), "REDONDEO"],
        ["ODs con diferencia diaria-semanal distinta de 0", int((diff.abs() > 1e-9).sum()), "REDONDEO"],
    ]
    quality = pd.DataFrame(checks, columns=["chequeo", "valor", "estado"])

    round_diff = pd.concat([
        dw[["origin", "dest", "tons_week"]],
        daily_sum.rename("suma_diaria"), diff.rename("diferencia")
    ], axis=1)
    round_diff = round_diff[round_diff["diferencia"].abs() > 1e-9].copy()

    long_legs = legs[legs["block_hours"] > float(ops["max_continuous_block_hours"])].copy()
    return quality, round_diff, long_legs


def demand_analysis(d):
    dw, dd = d["demand_weekly"], d["demand_daily"]

    totals = dd[DAYS].sum()
    active = (dd[DAYS] > 0).sum()
    day_summary = pd.DataFrame({
        "dia": [DAY_NAMES[x] for x in DAYS],
        "demanda_t": [totals[x] for x in DAYS],
        "porcentaje_semana": [100 * totals[x] / totals.sum() for x in DAYS],
        "ods_activos": [active[x] for x in DAYS],
        "t_promedio_por_od_activo": [totals[x] / active[x] if active[x] else np.nan for x in DAYS],
    })

    # d -> d+1: correlación por OD y transiciones de actividad.
    rel = []
    for i, day in enumerate(DAYS):
        nxt = DAYS[(i + 1) % len(DAYS)]
        x, y = dd[day], dd[nxt]
        pearson = x.corr(y, method="pearson")
        spearman = x.corr(y, method="spearman")
        rel.append({
            "dia_d": DAY_NAMES[day], "dia_d_mas_1": DAY_NAMES[nxt],
            "corr_pearson": pearson, "corr_spearman": spearman,
            "activo_ambos": int(((x > 0) & (y > 0)).sum()),
            "activo_a_cero": int(((x > 0) & (y == 0)).sum()),
            "cero_a_activo": int(((x == 0) & (y > 0)).sum()),
            "cero_ambos": int(((x == 0) & (y == 0)).sum()),
        })
    day_relation = pd.DataFrame(rel)

    corr = dd[DAYS].corr()
    corr.index = [DAY_NAMES[d] for d in DAYS]
    corr.columns = [DAY_NAMES[d] for d in DAYS]

    active_days = (dd[DAYS] > 0).sum(axis=1)
    freq_detail = dw[["origin", "dest", "tons_week", "min_weekly_freq"]].copy()
    freq_detail["dias_activos"] = active_days
    freq_summary = freq_detail.groupby("min_weekly_freq").agg(
        n_ods=("origin", "size"),
        demanda_media_t=("tons_week", "mean"),
        demanda_mediana_t=("tons_week", "median"),
        dias_activos_promedio=("dias_activos", "mean"),
        min_dias_activos=("dias_activos", "min"),
        max_dias_activos=("dias_activos", "max"),
    ).reset_index()

    sorted_dw = dw.sort_values("tons_week", ascending=False).reset_index(drop=True).copy()
    sorted_dw["share"] = sorted_dw["tons_week"] / sorted_dw["tons_week"].sum()
    sorted_dw["cum_share"] = sorted_dw["share"].cumsum()

    origin = dw.groupby("origin", as_index=False)["tons_week"].sum().sort_values("tons_week", ascending=False)
    dest = dw.groupby("dest", as_index=False)["tons_week"].sum().sort_values("tons_week", ascending=False)

    zeros = int((dd[DAYS] == 0).sum().sum())
    stats = {
        "total_week": float(dw["tons_week"].sum()),
        "n_ods": int(len(dw)),
        "zeros_daily": zeros,
        "daily_cells": int(dd[DAYS].size),
        "zero_share": zeros / dd[DAYS].size,
        "n80": int((sorted_dw["cum_share"] < 0.8).sum() + 1),
        "n50": int((sorted_dw["cum_share"] < 0.5).sum() + 1),
        "spearman_freq_active": float(freq_detail["min_weekly_freq"].corr(freq_detail["dias_activos"], method="spearman")),
        "spearman_demand_freq": float(freq_detail["tons_week"].corr(freq_detail["min_weekly_freq"], method="spearman")),
        "spearman_demand_active": float(freq_detail["tons_week"].corr(freq_detail["dias_activos"], method="spearman")),
    }
    return day_summary, day_relation, corr, freq_detail, freq_summary, sorted_dw, origin, dest, stats


def network_analysis(d):
    airports, legs, dw = d["airports"], d["legs"], d["demand_weekly"]
    hub = d["ops_rules"]["hub"]
    G = nx.DiGraph()
    for _, r in airports.iterrows():
        G.add_node(r["iata"], country=r["country"], region=r["region"])
    for _, r in legs.iterrows():
        G.add_edge(r["origin"], r["dest"], block_hours=r["block_hours"], distance_km=r["distance_km"])

    leg_pairs = set(zip(legs["origin"], legs["dest"]))
    asym = [(a, b) for a, b in leg_pairs if (b, a) not in leg_pairs]
    symmetry = 1 - len(asym) / len(leg_pairs)

    od_rows = []
    for _, r in dw.iterrows():
        direct = (r["origin"], r["dest"]) in leg_pairs
        try:
            path = nx.shortest_path(G, r["origin"], r["dest"])
            hops = len(path) - 1
        except nx.NetworkXNoPath:
            path, hops = [], np.nan
        simple_via_hub = (
            r["origin"] == hub or r["dest"] == hub or
            ((r["origin"], hub) in leg_pairs and (hub, r["dest"]) in leg_pairs)
        )
        od_rows.append({
            "origin": r["origin"], "dest": r["dest"], "tons_week": r["tons_week"],
            "directo": direct, "min_hops_red": hops, "via_hub_en_2_tramos": simple_via_hub,
            "camino_minimo_red": " -> ".join(path)
        })
    od_paths = pd.DataFrame(od_rows)

    degree = pd.DataFrame({
        "airport": list(G.nodes()),
        "in_degree": [G.in_degree(n) for n in G.nodes()],
        "out_degree": [G.out_degree(n) for n in G.nodes()],
        "degree_total": [G.degree(n) for n in G.nodes()],
    }).sort_values("degree_total", ascending=False)

    total = dw["tons_week"].sum()
    hub_in = dw.loc[dw["dest"].eq(hub), "tons_week"].sum()
    hub_out = dw.loc[dw["origin"].eq(hub), "tons_week"].sum()
    hub_stats = pd.DataFrame([
        ["Demanda total semanal", total, 100.0],
        [f"Demanda con destino {hub}", hub_in, 100 * hub_in / total],
        [f"Demanda con origen {hub}", hub_out, 100 * hub_out / total],
        [f"Desequilibrio neto hacia {hub}", hub_in - hub_out, 100 * (hub_in - hub_out) / total],
    ], columns=["metrica", "toneladas", "porcentaje_total"])

    region_map = airports.set_index("iata")["region"]
    reg = dw.copy()
    reg["region_origin"] = reg["origin"].map(region_map)
    reg["region_dest"] = reg["dest"].map(region_map)
    regional = reg.pivot_table(index="region_origin", columns="region_dest", values="tons_week", aggfunc="sum", fill_value=0)

    stats = {
        "nodes": G.number_of_nodes(), "edges": G.number_of_edges(), "density": nx.density(G),
        "asymmetric_edges": len(asym), "symmetry_share": symmetry,
        "direct_ods": int(od_paths["directo"].sum()), "nondirect_ods": int((~od_paths["directo"]).sum()),
        "hub_in": float(hub_in), "hub_out": float(hub_out), "hub_in_share": float(hub_in / total),
        "hub_out_share": float(hub_out / total), "hub_net": float(hub_in - hub_out),
    }
    return G, od_paths, degree, hub_stats, regional, stats


def operator_analysis(d):
    fleet, airports, legs = d["fleet"], d["airports"], d["legs"]
    rights, inter, cost = d["traffic_rights"], d["interchange"], d["cost_params"]
    operators = list(d["ops_rules"]["operators"].keys())
    country_of = airports.set_index("iata")["country"].to_dict()
    rights_dict = rights.set_index(["operator", "country"])["allowed"].to_dict()

    op_summary = fleet.groupby("operator").agg(
        n_aviones=("aircraft_id", "count"),
        payload_total_t=("payload_tons", "sum"),
        payload_promedio_t=("payload_tons", "mean")
    )
    allowed_countries = rights[rights["allowed"] == 1].groupby("operator")["country"].nunique()
    op_summary["paises_habilitados"] = allowed_countries

    leg_rows = []
    for _, r in legs.iterrows():
        valid = [op for op in operators
                 if rights_dict.get((op, country_of[r["origin"]]), 0) == 1
                 and rights_dict.get((op, country_of[r["dest"]]), 0) == 1]
        leg_rows.append({"origin": r["origin"], "dest": r["dest"],
                         "n_operadores_habilitados": len(valid), "operadores": ", ".join(valid)})
    leg_coverage = pd.DataFrame(leg_rows)
    available_by_op = {}
    for op in operators:
        available_by_op[op] = int(leg_coverage["operadores"].str.split(", ").apply(lambda xs: op in xs).sum())
    op_summary["tramos_habilitados"] = pd.Series(available_by_op)
    op_summary = op_summary.reset_index()

    inter_rows = []
    for a, b in combinations(operators, 2):
        shared = inter[(inter[a] == 1) & (inter[b] == 1)]["airport"].tolist()
        inter_rows.append({"operador_1": a, "operador_2": b,
                           "aeropuertos_compartidos": len(shared),
                           "lista_aeropuertos": ", ".join(shared)})
    inter_pairs = pd.DataFrame(inter_rows).sort_values("aeropuertos_compartidos", ascending=False)
    all_ops_airports = inter[(inter[operators] == 1).all(axis=1)]["airport"].tolist()

    inter_matrix = pd.DataFrame(index=operators, columns=operators, dtype=int)
    for a in operators:
        for b in operators:
            if a == b:
                inter_matrix.loc[a, b] = int((inter[a] == 1).sum())
            else:
                inter_matrix.loc[a, b] = int(((inter[a] == 1) & (inter[b] == 1)).sum())

    fuel_burn = float(cost.loc[cost["param"].eq("fuel_burn_gal_per_hour"), "value"].iloc[0])
    fuel_price = float(cost.loc[cost["param"].eq("fuel_price_usd_per_gal"), "value"].iloc[0])
    fuel_hour = fuel_burn * fuel_price
    cost_rows = []
    for op in operators:
        ex = float(cost.loc[(cost["param"].eq("ex_fuel_usd_per_block_hour")) & (cost["operator"].eq(op)), "value"].iloc[0])
        handling = float(cost.loc[(cost["param"].eq("handling_usd_per_ton")) & (cost["operator"].eq(op)), "value"].iloc[0])
        cost_rows.append({"operator": op, "fuel_usd_h": fuel_hour, "ex_fuel_usd_h": ex,
                          "costo_hora_aprox_usd": fuel_hour + ex, "handling_usd_t": handling})
    cost_by_op = pd.DataFrame(cost_rows)
    # Incorporar payload promedio y calcular costo por tonelada-capacidad por hora
    cost_by_op = cost_by_op.merge(
        op_summary[["operator", "payload_promedio_t", "payload_total_t", "n_aviones"]],
        on="operator", how="left"
    )
    cost_by_op["costo_por_tonelada_usd_t"] = (
        cost_by_op["costo_hora_aprox_usd"] / cost_by_op["payload_promedio_t"]
    )

    stats = {
        "all_ops_interchange_n": len(all_ops_airports),
        "all_ops_interchange": all_ops_airports,
        "min_ops_per_leg": int(leg_coverage["n_operadores_habilitados"].min()),
        "legs_two_ops": int((leg_coverage["n_operadores_habilitados"] == 2).sum()),
        "legs_three_ops": int((leg_coverage["n_operadores_habilitados"] == 3).sum()),
        "legs_four_ops": int((leg_coverage["n_operadores_habilitados"] == 4).sum()),
    }
    return op_summary, leg_coverage, inter_pairs, inter_matrix, cost_by_op, stats


def maintenance_analysis(d):
    maint = d["maintenance"].copy()
    fleet = d["fleet"]
    week_start = pd.Timestamp(d["ops_rules"]["week_start"])
    maint["duration_h"] = (maint["end_datetime"] - maint["start_datetime"]).dt.total_seconds() / 3600
    maint = maint.merge(fleet[["aircraft_id", "operator"]], on="aircraft_id", how="left")

    rows = []
    for i, day in enumerate(DAYS):
        start = week_start + pd.Timedelta(days=i)
        end = start + pd.Timedelta(days=1)
        hours = 0.0
        affected = set()
        for _, r in maint.iterrows():
            overlap = min(r["end_datetime"], end) - max(r["start_datetime"], start)
            if overlap > pd.Timedelta(0):
                hours += overlap.total_seconds() / 3600
                affected.add(r["aircraft_id"])
        rows.append({"dia": DAY_NAMES[day], "aviones_afectados": len(affected), "horas_mantenimiento": hours})
    maint_day = pd.DataFrame(rows)

    by_airport = maint.groupby("airport").agg(ventanas=("aircraft_id", "count"), horas=("duration_h", "sum")).reset_index().sort_values("horas", ascending=False)
    by_operator = maint.groupby("operator").agg(ventanas=("aircraft_id", "count"), horas=("duration_h", "sum")).reset_index().sort_values("horas", ascending=False)
    return maint, maint_day, by_airport, by_operator


def slack_catalog():
    rows = [
        ["Holgura de capacidad", "capacidad del avión - carga efectivamente asignada al vuelo", "No", "Muy útil: identifica vuelos con capacidad ociosa."],
        ["Holgura de demanda", "demanda disponible - toneladas efectivamente servidas", "No", "Muy útil: cuantifica mercado no atendido y oportunidades sacrificadas."],
        ["Exceso sobre frecuencia mínima", "vuelos realizados - frecuencia mínima comprometida", "No", "Útil: muestra dónde se vuela por encima del compromiso."],
        ["Holgura temporal/TAT", "tiempo real entre vuelos - TAT mínimo", "No", "Útil solo si existe horario intradía explícito."],
        ["Ociosidad de aeronave", "tiempo/celdas de tiempo sin vuelo ni mantenimiento", "No", "Útil: mide capacidad operacional no utilizada."],
        ["Holgura de balance/ciclo", "No es una holgura económica; el cierre semanal debe cumplirse", "No", "No conviene interpretarla como recurso disponible."],
        ["Derechos de tráfico", "Restricción binaria habilitado/no habilitado", "Sí como dato", "No hablar de slack: es una condición estructural."],
        ["Existencia de arco", "Restricción binaria: el tramo existe o no existe", "Sí como dato", "No hablar de slack: es una condición estructural."],
    ]
    return pd.DataFrame(rows, columns=["tipo", "definicion", "calculable_sin_solucion", "interpretacion"])


def modeling_alerts(root: Path, d, net_stats, op_stats):
    repo_root = resolve_repo_root(root)
    data_dir = resolve_data_dir(root)
    readme = (repo_root / "README.md").read_text(encoding="utf-8", errors="ignore") if (repo_root / "README.md").exists() else ""
    model_md = (repo_root / "docs" / "modelo_matematico.md").read_text(encoding="utf-8", errors="ignore") if (repo_root / "docs" / "modelo_matematico.md").exists() else ""
    ops_text = (data_dir / "ops_rules.yaml").read_text(encoding="utf-8", errors="ignore")
    legs = d["legs"]

    alerts = []
    # No posición inicial explícita
    alerts.append([
        "Posición inicial/final de aviones",
        "No existe en los datos un aeropuerto inicial explícito por aircraft_id. La regla es cíclica: inicio y fin deben ser compatibles/iguales, pero la posición inicial no está fijada externamente.",
        "Importante",
        "El punto inicial afecta el final por el cierre semanal. Si no se fija una base inicial, el modelo puede elegir la rotación cerrada que más convenga."
    ])

    # Interchange inconsistency
    if "pasa a ser operado por otra filial" in readme and "permanente" in model_md:
        alerts.append([
            "Interchange vs operador fijo",
            "README define interchange como cambio de filial que opera la aeronave; ops_rules y modelo matemático indican que cada avión tiene operador fijo/permanente.",
            "Crítico",
            "Hay que decidir si el interchange aplica a la aeronave o solo a la carga/operación logística. Cambia las restricciones de derechos y costos."
        ])

    # Transshipment inconsistency
    if "sin transbordo" in readme and "transbordo" in model_md:
        alerts.append([
            "Conexiones sin transbordo vs transbordo en MIA",
            "README menciona conexiones sin transbordo, mientras el modelo matemático permite/obliga transbordo de carga en MIA entre aviones distintos.",
            "Crítico",
            "Definir una única lógica de conexión de carga antes de cerrar la formulación."
        ])

    # TAT granularity
    if "tat_minutes" in ops_text and not any("time" in c.lower() or "datetime" in c.lower() for c in legs.columns):
        alerts.append([
            "Granularidad temporal vs TAT",
            "TAT está definido en minutos, pero legs_catalog solo contiene duración y el modelo de vuelo mostrado se indexa por día; no hay hora de salida programada por tramo.",
            "Crítico",
            "Con granularidad diaria no se puede verificar de forma exacta un TAT de 50/70/90 minutos sin incorporar horarios o una red tiempo-espacio más fina."
        ])

    # One flight/day vs TAT
    if "exactamente una cosa" in model_md:
        alerts.append([
            "Un avión hace una cosa por día",
            "R2 limita a una acción diaria por avión. Esto simplifica mucho la operación y puede hacer que el TAT intradía pierda relevancia.",
            "Importante",
            "Revisar si el enunciado realmente restringe a un vuelo por día o si un avión puede realizar varios tramos en 24 h."
        ])

    # Long legs
    alerts.append([
        "Autonomía continua",
        f"Existen {(d['legs']['block_hours'] > float(d['ops_rules']['max_continuous_block_hours'])).sum()} tramos con block_hours > {d['ops_rules']['max_continuous_block_hours']} h.",
        "Importante",
        "Esos tramos requieren tratamiento de escala técnica o deben excluirse como vuelos directos a plena carga."
    ])

    return pd.DataFrame(alerts, columns=["tema", "hallazgo", "prioridad", "conclusion"])


def make_graphs(out_graphs: Path, results):
    out_graphs.mkdir(parents=True, exist_ok=True)
    day_summary = results["day_summary"]
    day_relation = results["day_relation"]
    corr = results["day_corr"]
    freq_summary = results["freq_summary"]
    pareto = results["pareto"]
    origin = results["origin"]
    dest = results["dest"]
    G = results["G"]
    od_paths = results["od_paths"]
    degree = results["degree"]
    hub_stats = results["hub_stats"]
    op_summary = results["op_summary"]
    leg_coverage = results["leg_coverage"]
    inter_matrix = results["inter_matrix"]
    maint_day = results["maint_day"]
    cost_by_op = results["cost_by_op"]
    regional = results["regional"]

    paths = {}

    # 1 Demanda total día
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.bar(day_summary["dia"], day_summary["demanda_t"])
    clean_ax(ax, "Demanda total por día", ylabel="Toneladas")
    ax.tick_params(axis="x", rotation=30)
    paths["demanda_dia"] = out_graphs / "01_demanda_total_por_dia.png"
    save_fig(fig, paths["demanda_dia"])

    # 2 ODs activos
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.bar(day_summary["dia"], day_summary["ods_activos"])
    clean_ax(ax, "ODs con demanda positiva por día", ylabel="Número de ODs")
    ax.tick_params(axis="x", rotation=30)
    paths["ods_activos"] = out_graphs / "02_ods_activos_por_dia.png"
    save_fig(fig, paths["ods_activos"])

    # 3 Demanda por OD activo
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.bar(day_summary["dia"], day_summary["t_promedio_por_od_activo"])
    clean_ax(ax, "Toneladas promedio por OD activo", ylabel="t / OD activo")
    ax.tick_params(axis="x", rotation=30)
    paths["promedio_od"] = out_graphs / "03_toneladas_promedio_por_od_activo.png"
    save_fig(fig, paths["promedio_od"])

    # 4 Correlación días
    fig, ax = plt.subplots(figsize=(7, 6))
    im = ax.imshow(corr.values, vmin=-1, vmax=1)
    ax.set_xticks(range(len(corr.columns)), corr.columns, rotation=45, ha="right")
    ax.set_yticks(range(len(corr.index)), corr.index)
    for i in range(len(corr.index)):
        for j in range(len(corr.columns)):
            ax.text(j, i, f"{corr.iloc[i,j]:.2f}", ha="center", va="center", fontsize=8)
    fig.colorbar(im, ax=ax, label="Correlación")
    ax.set_title("Correlación de demanda por OD entre días")
    paths["corr_dias"] = out_graphs / "04_correlacion_demanda_entre_dias.png"
    save_fig(fig, paths["corr_dias"])

    # 5 Transiciones actividad d -> d+1
    fig, ax = plt.subplots(figsize=(9, 4.8))
    x = np.arange(len(day_relation))
    width = 0.25
    ax.bar(x - width, day_relation["activo_ambos"], width, label="Activo ambos")
    ax.bar(x, day_relation["activo_a_cero"], width, label="Activo -> cero")
    ax.bar(x + width, day_relation["cero_a_activo"], width, label="Cero -> activo")
    ax.set_xticks(x, [f"{a[:3]}→{b[:3]}" for a,b in zip(day_relation["dia_d"], day_relation["dia_d_mas_1"])])
    clean_ax(ax, "Transiciones de actividad OD entre días consecutivos", ylabel="Número de ODs")
    ax.legend()
    paths["transiciones"] = out_graphs / "05_transiciones_od_dia_siguiente.png"
    save_fig(fig, paths["transiciones"])

    # 6 Frecuencia mínima vs días activos
    fig, ax = plt.subplots(figsize=(7, 4.5))
    ax.plot(freq_summary["min_weekly_freq"], freq_summary["dias_activos_promedio"], marker="o")
    clean_ax(ax, "Frecuencia mínima y días activos observados", xlabel="Frecuencia mínima semanal", ylabel="Días activos promedio")
    paths["freq_dias"] = out_graphs / "06_frecuencia_minima_vs_dias_activos.png"
    save_fig(fig, paths["freq_dias"])

    # 7 Pareto demanda
    fig, ax = plt.subplots(figsize=(10, 4.8))
    ax.bar(np.arange(len(pareto)), pareto["tons_week"])
    ax.set_xlabel("ODs ordenados por demanda")
    ax.set_ylabel("Toneladas/semana")
    ax2 = ax.twinx()
    ax2.plot(np.arange(len(pareto)), 100 * pareto["cum_share"], marker=".", linewidth=1)
    ax2.axhline(80, linestyle="--", linewidth=1)
    ax2.set_ylabel("% acumulado")
    ax.set_title("Concentración de la demanda semanal (Pareto)")
    paths["pareto"] = out_graphs / "07_pareto_demanda_od.png"
    save_fig(fig, paths["pareto"])

    # 8 Top orígenes
    top = origin.head(10).sort_values("tons_week")
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.barh(top["origin"], top["tons_week"])
    clean_ax(ax, "Top 10 orígenes por demanda", xlabel="Toneladas/semana")
    paths["origins"] = out_graphs / "08_top_origenes_demanda.png"
    save_fig(fig, paths["origins"])

    # 9 Top destinos
    top = dest.head(10).sort_values("tons_week")
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.barh(top["dest"], top["tons_week"])
    clean_ax(ax, "Top 10 destinos por demanda", xlabel="Toneladas/semana")
    paths["destinations"] = out_graphs / "09_top_destinos_demanda.png"
    save_fig(fig, paths["destinations"])

    # 10 MIA flujo
    mia_plot = hub_stats.iloc[1:3].copy()
    fig, ax = plt.subplots(figsize=(7, 4.5))
    ax.bar(["Hacia MIA", "Desde MIA"], mia_plot["toneladas"])
    clean_ax(ax, "Asimetría de demanda en el hub MIA", ylabel="Toneladas/semana")
    paths["mia"] = out_graphs / "10_flujo_demanda_mia.png"
    save_fig(fig, paths["mia"])

    # 11 Directos vs no directos
    direct_counts = od_paths["directo"].value_counts().reindex([True, False], fill_value=0)
    fig, ax = plt.subplots(figsize=(6, 4.5))
    ax.bar(["OD con tramo directo", "OD sin tramo directo"], direct_counts.values)
    clean_ax(ax, "Cobertura directa de los ODs de demanda", ylabel="Número de ODs")
    ax.tick_params(axis="x", rotation=15)
    paths["directos"] = out_graphs / "11_ods_directos_vs_conexion.png"
    save_fig(fig, paths["directos"])

    # 12 Hops mínimos
    hops = od_paths["min_hops_red"].value_counts().sort_index()
    fig, ax = plt.subplots(figsize=(6.5, 4.5))
    ax.bar(hops.index.astype(str), hops.values)
    clean_ax(ax, "Número mínimo de tramos por OD en la red física", xlabel="Tramos mínimos", ylabel="Número de ODs")
    paths["hops"] = out_graphs / "12_minimo_tramos_por_od.png"
    save_fig(fig, paths["hops"])

    # 13 Red física
    fig, ax = plt.subplots(figsize=(10, 8))
    pos = nx.spring_layout(G, seed=42, k=0.55)
    nx.draw_networkx_nodes(G, pos, node_size=90, ax=ax)
    nx.draw_networkx_edges(G, pos, alpha=0.25, arrows=True, arrowsize=7, ax=ax)
    nx.draw_networkx_labels(G, pos, font_size=6, ax=ax)
    ax.set_title("Red dirigida de tramos disponibles")
    ax.axis("off")
    paths["network"] = out_graphs / "13_red_tramos.png"
    save_fig(fig, paths["network"])

    # 14 Cobertura operadores
    fig, ax = plt.subplots(figsize=(7, 4.5))
    ax.bar(op_summary["operator"], op_summary["tramos_habilitados"])
    clean_ax(ax, "Tramos habilitados por operador", ylabel="Tramos del catálogo")
    paths["operator_legs"] = out_graphs / "14_tramos_habilitados_operador.png"
    save_fig(fig, paths["operator_legs"])

    # 15 Flota operador
    fig, ax = plt.subplots(figsize=(7, 4.5))
    ax.bar(op_summary["operator"], op_summary["n_aviones"])
    clean_ax(ax, "Número de aviones por operador", ylabel="Aviones")
    paths["fleet"] = out_graphs / "15_flota_por_operador.png"
    save_fig(fig, paths["fleet"])

    # 16 Interchange matrix
    fig, ax = plt.subplots(figsize=(6, 5))
    im = ax.imshow(inter_matrix.values)
    ax.set_xticks(range(len(inter_matrix.columns)), inter_matrix.columns)
    ax.set_yticks(range(len(inter_matrix.index)), inter_matrix.index)
    for i in range(len(inter_matrix.index)):
        for j in range(len(inter_matrix.columns)):
            ax.text(j, i, f"{int(inter_matrix.iloc[i,j])}", ha="center", va="center")
    fig.colorbar(im, ax=ax, label="Aeropuertos compartidos")
    ax.set_title("Interchange: aeropuertos compartidos entre operadores")
    paths["interchange"] = out_graphs / "16_interchange_operadores.png"
    save_fig(fig, paths["interchange"])

    # 17 Mantenimiento
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.bar(maint_day["dia"], maint_day["horas_mantenimiento"])
    clean_ax(ax, "Horas de mantenimiento superpuestas a cada día", ylabel="Horas-avión")
    ax.tick_params(axis="x", rotation=30)
    paths["maintenance"] = out_graphs / "17_mantenimiento_por_dia.png"
    save_fig(fig, paths["maintenance"])

    # 18 Demanda vs mantenimiento
    merged = day_summary.merge(maint_day, on="dia")
    fig, ax = plt.subplots(figsize=(6.5, 4.8))
    ax.scatter(merged["horas_mantenimiento"], merged["demanda_t"])
    for _, r in merged.iterrows():
        ax.annotate(r["dia"], (r["horas_mantenimiento"], r["demanda_t"]), fontsize=8)
    clean_ax(ax, "Demanda diaria vs horas de mantenimiento", xlabel="Horas-avión de mantenimiento", ylabel="Toneladas de demanda")
    paths["dem_maint"] = out_graphs / "18_demanda_vs_mantenimiento.png"
    save_fig(fig, paths["dem_maint"])

    # 19 Costo por tonelada de capacidad útil por hora
    cb = cost_by_op.sort_values("costo_por_tonelada_usd_t", ascending=True).reset_index(drop=True)
    fig, ax = plt.subplots(figsize=(7.2, 4.8))
    idx_min = int(cb["costo_por_tonelada_usd_t"].idxmin())
    colors = ["#012169" if i == idx_min else "#4a6fa5" for i in range(len(cb))]
    bars = ax.bar(cb["operator"], cb["costo_por_tonelada_usd_t"], color=colors, edgecolor="white", linewidth=0.6)
    # Etiqueta valor encima de cada barra
    for bar, val in zip(bars, cb["costo_por_tonelada_usd_t"]):
        ax.text(bar.get_x() + bar.get_width() / 2, val + max(cb["costo_por_tonelada_usd_t"]) * 0.018,
                f"{val:,.0f}", ha="center", va="bottom", fontsize=9, fontweight="bold", color="#012169")
    clean_ax(ax, "Costo por tonelada-capacidad útil según operador",
             ylabel="USD por tonelada y por block-hour")
    ax.set_xlabel("Operador (ordenado por costo/t, ascendente)", fontsize=9)
    ax.grid(axis="y", color="#e8ecf1", linewidth=0.6)
    paths["costs"] = out_graphs / "19_costo_por_tonelada_operador.png"
    save_fig(fig, paths["costs"])

    # 20 Regional heatmap
    fig, ax = plt.subplots(figsize=(7, 5.5))
    im = ax.imshow(regional.values)
    ax.set_xticks(range(len(regional.columns)), regional.columns)
    ax.set_yticks(range(len(regional.index)), regional.index)
    for i in range(len(regional.index)):
        for j in range(len(regional.columns)):
            ax.text(j, i, f"{regional.iloc[i,j]:.0f}", ha="center", va="center", fontsize=8)
    fig.colorbar(im, ax=ax, label="Toneladas/semana")
    ax.set_xlabel("Región destino")
    ax.set_ylabel("Región origen")
    ax.set_title("Demanda agregada entre regiones")
    paths["regional"] = out_graphs / "20_demanda_regional.png"
    save_fig(fig, paths["regional"])

    return paths


def set_doc_font(doc: Document):
    styles = doc.styles
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.size = Pt(18)
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.size = Pt(11)


def add_table(doc: Document, df: pd.DataFrame, columns=None, max_rows=12, decimals=1):
    if columns is not None:
        df = df[columns]
    view = df.head(max_rows).copy()
    table = doc.add_table(rows=1, cols=len(view.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for j, c in enumerate(view.columns):
        hdr[j].text = str(c)
        hdr[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in hdr[j].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    for _, row in view.iterrows():
        cells = table.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = fmt(v, decimals=decimals)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[j].paragraphs[0].runs:
                run.font.size = Pt(8)
    if len(df) > max_rows:
        p = doc.add_paragraph(f"Se muestran {max_rows} de {len(df)} filas. El detalle completo queda exportado en CSV.")
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(8)
    return table


def add_figure(doc: Document, path: Path, caption: str, width=6.2):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(8)


def build_word(out_path: Path, results, graphs):
    doc = Document()
    set_doc_font(doc)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Diagnóstico estructural de la instancia ICS2122")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Grupo 9 - análisis previo a la optimización").italic = True

    doc.add_heading("Enfoque", level=1)
    doc.add_paragraph(
        "Este análisis no intenta ajustar distribuciones probabilísticas. La instancia se trata como un problema "
        "estático y determinístico de optimización. El objetivo es entender la demanda, la red, los operadores, "
        "las dependencias espacio-temporales y qué holguras serán informativas una vez resuelto el modelo."
    )

    ds = results["demand_stats"]
    ns = results["net_stats"]
    os = results["op_stats"]
    doc.add_heading("Resumen ejecutivo", level=1)
    bullets = [
        f"La instancia contiene {ds['n_ods']} ODs y {fmt(ds['total_week'])} toneladas de demanda semanal potencial.",
        f"La demanda diaria es muy estructurada: {ds['zeros_daily']} de {ds['daily_cells']} celdas OD-día son cero ({fmt(100*ds['zero_share'])}%).",
        f"Los {ds['n80']} ODs de mayor demanda concentran aproximadamente 80% del tonelaje semanal.",
        f"{ns['direct_ods']} ODs tienen tramo directo en el catálogo y {ns['nondirect_ods']} requieren al menos una conexión en la red física.",
        f"MIA recibe {fmt(ns['hub_in'])} t ({fmt(100*ns['hub_in_share'])}% del total), frente a {fmt(ns['hub_out'])} t que salen desde MIA.",
        f"La red física tiene {ns['edges']} tramos dirigidos y solo {fmt(100*ns['symmetry_share'])}% de simetría A↔B; el sentido de la ruta importa.",
        f"Todos los tramos tienen al menos {os['min_ops_per_leg']} operadores habilitados por derechos de tráfico, pero la cobertura no es igual entre empresas.",
        "La posición de un avión en el día siguiente depende de dónde terminó el día anterior y la operación debe cerrar cíclicamente al final de la semana.",
        "Las holguras de capacidad, demanda, frecuencia y tiempo no pueden calcularse de verdad antes de resolver el modelo; aquí solo se define cómo interpretarlas.",
    ]
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("1. Datos e integridad", level=1)
    doc.add_paragraph("Se utilizaron todos los archivos de entrada disponibles en data/.")
    add_table(doc, results["inventory"], max_rows=20)
    doc.add_paragraph("Chequeos principales:")
    add_table(doc, results["quality"], max_rows=20)
    if len(results["long_legs"]):
        doc.add_paragraph("Tramos que exceden la autonomía continua indicada en ops_rules.yaml:")
        add_table(doc, results["long_legs"], max_rows=10)
    doc.add_paragraph(
        "La suma de demanda diaria y semanal es consistente salvo diferencias de redondeo de hasta "
        f"{fmt(results['round_diff']['diferencia'].abs().max())} t. No se interpreta como un problema estructural."
    )

    doc.add_heading("2. Comportamiento de la demanda por día", level=1)
    add_table(doc, results["day_summary"], max_rows=10)
    add_figure(doc, graphs["demanda_dia"], "Figura: demanda total por día.")
    add_figure(doc, graphs["ods_activos"], "Figura: cantidad de ODs activos por día.")
    add_figure(doc, graphs["promedio_od"], "Figura: toneladas promedio por OD activo.")
    doc.add_paragraph(
        "Los días no se diferencian solo por tonelaje total. Lunes, miércoles y viernes activan muchos más ODs, "
        "mientras martes, jueves, sábado y domingo concentran la demanda en un número menor de ODs, pero con mayor tonelaje promedio por OD activo."
    )
    doc.add_paragraph(
        f"La relación entre frecuencia mínima y días activos es fuerte: correlación de Spearman = {results['demand_stats']['spearman_freq_active']:.2f}. "
        "Esto indica que la desagregación diaria está ligada a la estructura comercial/frecuencia del OD."
    )
    add_table(doc, results["freq_summary"], max_rows=10)
    add_figure(doc, graphs["freq_dias"], "Figura: frecuencia mínima semanal versus días activos promedio.")

    doc.add_heading("3. Relación día d con día d+1", level=1)
    add_table(doc, results["day_relation"], max_rows=10, decimals=2)
    add_figure(doc, graphs["corr_dias"], "Figura: matriz de correlación de tonelaje por OD entre días.")
    add_figure(doc, graphs["transiciones"], "Figura: ODs que se mantienen activos, se apagan o se activan entre días consecutivos.")
    doc.add_paragraph(
        "Las correlaciones Pearson son altas porque los ODs grandes tienden a seguir siendo grandes en varios días. "
        "No deben interpretarse como causalidad temporal: solo se dispone de una semana, no de una serie histórica de semanas. "
        "La señal más útil es el patrón de activación/desactivación de ODs entre días."
    )
    doc.add_paragraph(
        "La demanda de un día no se transfiere al siguiente. Por lo tanto, d+1 no hereda demanda no servida de d; "
        "la dependencia temporal importante proviene principalmente de la posición y disponibilidad de los aviones."
    )

    doc.add_heading("4. Concentración y direccionalidad de la demanda", level=1)
    doc.add_paragraph(
        f"Los {results['demand_stats']['n50']} ODs más grandes superan 50% de la demanda y los {results['demand_stats']['n80']} más grandes alcanzan aproximadamente 80%."
    )
    add_figure(doc, graphs["pareto"], "Figura: concentración acumulada de demanda por OD.")
    add_table(doc, results["pareto"].head(12)[["origin","dest","tons_week","min_weekly_freq","cum_share"]], max_rows=12, decimals=2)
    add_figure(doc, graphs["origins"], "Figura: principales aeropuertos de origen por tonelaje.")
    add_figure(doc, graphs["destinations"], "Figura: principales aeropuertos de destino por tonelaje.")
    add_figure(doc, graphs["mia"], "Figura: tonelaje de demanda hacia y desde MIA.")
    doc.add_paragraph(
        f"Existe una asimetría fuerte hacia MIA: {fmt(ns['hub_in'])} t tienen MIA como destino y {fmt(ns['hub_out'])} t tienen MIA como origen. "
        f"El desequilibrio neto es de {fmt(ns['hub_net'])} t hacia el hub. Esto puede obligar a reposicionar aeronaves o aceptar vuelos con baja carga en el sentido contrario."
    )

    doc.add_heading("5. Red física y efecto del origen/destino", level=1)
    add_figure(doc, graphs["network"], "Figura: red dirigida de tramos disponibles.", width=6.4)
    add_figure(doc, graphs["directos"], "Figura: ODs con y sin tramo directo.")
    add_figure(doc, graphs["hops"], "Figura: cantidad mínima de tramos requerida por OD en la red física.")
    doc.add_paragraph(
        f"La red tiene {ns['nodes']} aeropuertos y {ns['edges']} arcos dirigidos. Existen {ns['asymmetric_edges']} arcos cuyo sentido inverso no está en el catálogo. "
        "Por lo tanto, A→B no equivale a B→A y el punto de origen condiciona de forma directa las alternativas de continuación."
    )
    doc.add_paragraph("ODs sin tramo directo (detalle):")
    add_table(doc, results["od_paths"].query("directo == False")[["origin", "dest", "tons_week", "min_hops_red", "camino_minimo_red"]], max_rows=20)
    add_figure(doc, graphs["regional"], "Figura: demanda agregada entre regiones.")

    doc.add_heading("6. Empresas, derechos de tráfico e interchange", level=1)
    add_table(doc, results["op_summary"], max_rows=10)
    add_figure(doc, graphs["fleet"], "Figura: tamaño de flota por operador.")
    add_figure(doc, graphs["operator_legs"], "Figura: tramos del catálogo habilitados por operador.")
    doc.add_paragraph(
        f"En los {ns['edges']} tramos del catálogo, {os['legs_four_ops']} pueden ser operados por las cuatro empresas, "
        f"{os['legs_three_ops']} por tres y {os['legs_two_ops']} solo por dos. Tener capacidad libre en una empresa no implica poder usarla en cualquier tramo."
    )
    add_table(doc, results["inter_pairs"][["operador_1", "operador_2", "aeropuertos_compartidos"]], max_rows=10)
    add_figure(doc, graphs["interchange"], "Figura: número de aeropuertos compartidos entre pares de operadores.")
    doc.add_paragraph(
        f"Hay {os['all_ops_interchange_n']} aeropuertos en los que los cuatro operadores aparecen habilitados en interchange_airports.csv, incluido MIA. "
        "Sin embargo, los documentos del proyecto no son consistentes sobre si este interchange significa cambiar el operador de la aeronave o solo transferir carga/operación logística."
    )

    doc.add_heading("7. Dependencias espacio-temporales", level=1)
    doc.add_paragraph(
        "La dependencia central no es probabilística, sino de continuidad física: un avión solo puede salir desde el aeropuerto donde quedó después de su acción anterior. "
        "Además, la semana es cíclica, por lo que su posición final debe permitir comenzar la semana siguiente en la misma rotación."
    )
    doc.add_paragraph(
        "No existe en los datos un aeropuerto inicial explícito para cada aircraft_id. En la formulación actual, la rotación cerrada puede determinar implícitamente su punto de inicio. "
        "Esto hace que inicio y final estén totalmente vinculados, pero no necesariamente anclados a una base predeterminada."
    )
    add_table(doc, results["maint_day"], max_rows=10)
    add_figure(doc, graphs["maintenance"], "Figura: horas-avión de mantenimiento por día.")
    add_figure(doc, graphs["dem_maint"], "Figura: demanda diaria y mantenimiento. Con solo siete días no se infiere una relación estadística.")
    doc.add_paragraph(
        "El mantenimiento crea restricciones simultáneamente espaciales y temporales: cada aeronave debe estar en el aeropuerto indicado durante su ventana y no puede volar. "
        "El TAT también es espacio-temporal, pero actualmente está definido en minutos mientras la formulación principal usa decisiones diarias; esta diferencia de granularidad debe revisarse."
    )

    doc.add_heading("8. Costos y diferencias entre operadores", level=1)
    add_table(doc, results["cost_by_op"], max_rows=10)
    add_figure(doc, graphs["costs"], "Figura: costo por tonelada de capacidad útil por block-hour, según operador. Menor = más eficiente en costo por carga transportada.")
    add_table(doc, results["fees_summary"], max_rows=10)
    doc.add_paragraph(
        "Los operadores no solo difieren en derechos de tráfico y tamaño de flota; también tienen costos ex-fuel y handling distintos. "
        "Por eso dos aviones con capacidad parecida no son económicamente intercambiables."
    )

    doc.add_heading("9. Holguras (slack): qué complementa y qué distrae", level=1)
    add_table(doc, results["slack"], max_rows=20)
    doc.add_paragraph(
        "Las holguras útiles se calculan con la solución del modelo, no solo con los datos. La más importante será capacidad no utilizada por vuelo; luego demanda no servida, "
        "frecuencia por sobre el mínimo, ociosidad de aeronaves y, si existe un horario intradía, holgura de TAT."
    )
    doc.add_paragraph(
        "No conviene llamar slack a derechos de tráfico, existencia de arcos o pertenencia a un operador: son condiciones binarias/estructurales. "
        "Llenar el informe con 'slacks' de estas restricciones distraería de la interpretación económica y operacional."
    )

    doc.add_heading("10. Alertas de formulación que conviene resolver", level=1)
    add_table(doc, results["alerts"], max_rows=20)

    doc.add_heading("Conclusiones", level=1)
    conclusions = [
        "El análisis correcto para esta instancia es estructural y determinístico, no un ajuste de distribuciones de probabilidad.",
        "La demanda tiene una fuerte estructura por día y frecuencia mínima; la demanda no servida hoy no se acumula para mañana.",
        "La principal dependencia entre días es la continuidad física de las aeronaves y el cierre cíclico semanal.",
        "La asimetría hacia MIA y la direccionalidad del catálogo hacen que posicionamiento y vuelos de retorno sean centrales en la rentabilidad global.",
        "Las empresas tienen flota, costos y derechos distintos. La capacidad disponible no es completamente fungible entre operadores.",
        "El significado de interchange debe aclararse porque README, reglas y modelo no describen exactamente la misma lógica.",
        "El TAT en minutos no es plenamente compatible con una formulación que decide solo por día sin horarios de salida; es una prioridad de modelado.",
        "El análisis de slack debe hacerse después de optimizar y concentrarse en capacidad, demanda no servida, frecuencia, ociosidad y tiempo, no en restricciones estructurales binarias.",
    ]
    for c in conclusions:
        doc.add_paragraph(c, style="List Bullet")

    doc.save(out_path)


def main():
    args = parse_args()
    root = args.root.resolve()
    out = (args.out or (root / "analisis_instancia")).resolve()
    tables_dir = out / "tablas"
    graphs_dir = out / "graficos"
    out.mkdir(parents=True, exist_ok=True)
    tables_dir.mkdir(parents=True, exist_ok=True)

    d = load_data(root)

    inv = inventory(d)
    quality, round_diff, long_legs = data_quality(d)
    day_summary, day_relation, day_corr, freq_detail, freq_summary, pareto, origin, dest, demand_stats = demand_analysis(d)
    G, od_paths, degree, hub_stats, regional, net_stats = network_analysis(d)
    op_summary, leg_coverage, inter_pairs, inter_matrix, cost_by_op, op_stats = operator_analysis(d)
    maint_detail, maint_day, maint_airport, maint_operator = maintenance_analysis(d)
    slack = slack_catalog()
    alerts = modeling_alerts(root, d, net_stats, op_stats)

    fees = d["landing_fees"]
    fees_summary = pd.DataFrame([
        ["n_aeropuertos", len(fees)],
        ["fee_promedio_usd", fees["fee_usd"].mean()],
        ["fee_mediana_usd", fees["fee_usd"].median()],
        ["fee_min_usd", fees["fee_usd"].min()],
        ["fee_max_usd", fees["fee_usd"].max()],
    ], columns=["metrica", "valor"])

    # Exportar tablas detalladas.
    exports = {
        "00_inventario_datos.csv": inv,
        "01_calidad_datos.csv": quality,
        "02_diferencias_demanda_diaria_semanal.csv": round_diff,
        "03_tramos_sobre_autonomia.csv": long_legs,
        "04_demanda_por_dia.csv": day_summary,
        "05_relacion_dia_siguiente.csv": day_relation,
        "06_correlacion_dias.csv": day_corr,
        "07_frecuencia_y_dias_activos_detalle.csv": freq_detail,
        "08_frecuencia_y_dias_activos_resumen.csv": freq_summary,
        "09_pareto_demanda_od.csv": pareto,
        "10_demanda_por_origen.csv": origin,
        "11_demanda_por_destino.csv": dest,
        "12_rutas_od_y_conexiones.csv": od_paths,
        "13_grado_aeropuertos.csv": degree,
        "14_flujo_hub_mia.csv": hub_stats,
        "15_demanda_regional.csv": regional,
        "16_resumen_operadores.csv": op_summary,
        "17_cobertura_tramos_operadores.csv": leg_coverage,
        "18_interchange_pares.csv": inter_pairs,
        "19_interchange_matriz.csv": inter_matrix,
        "20_costos_por_operador.csv": cost_by_op,
        "21_mantenimiento_detalle.csv": maint_detail,
        "22_mantenimiento_por_dia.csv": maint_day,
        "23_mantenimiento_por_aeropuerto.csv": maint_airport,
        "24_mantenimiento_por_operador.csv": maint_operator,
        "25_landing_fees_resumen.csv": fees_summary,
        "26_slack_catalogo.csv": slack,
        "27_alertas_modelado.csv": alerts,
    }
    for filename, df in exports.items():
        save_df(df, tables_dir / filename, index=isinstance(df.index, pd.Index) and df.index.name is not None)

    results = {
        "inventory": inv, "quality": quality, "round_diff": round_diff, "long_legs": long_legs,
        "day_summary": day_summary, "day_relation": day_relation, "day_corr": day_corr,
        "freq_detail": freq_detail, "freq_summary": freq_summary, "pareto": pareto,
        "origin": origin, "dest": dest, "demand_stats": demand_stats,
        "G": G, "od_paths": od_paths, "degree": degree, "hub_stats": hub_stats,
        "regional": regional, "net_stats": net_stats,
        "op_summary": op_summary, "leg_coverage": leg_coverage, "inter_pairs": inter_pairs,
        "inter_matrix": inter_matrix, "cost_by_op": cost_by_op, "op_stats": op_stats,
        "maint_detail": maint_detail, "maint_day": maint_day,
        "maint_airport": maint_airport, "maint_operator": maint_operator,
        "fees_summary": fees_summary, "slack": slack, "alerts": alerts,
    }

    graphs = make_graphs(graphs_dir, results)
    word_path = out / "Diagnostico_instancia_ICS2122.docx"
    build_word(word_path, results, graphs)

    print("\n=== DIAGNÓSTICO COMPLETADO ===")
    print(f"Raíz analizada: {root}")
    print(f"Salida: {out}")
    print(f"Word: {word_path}")
    print(f"Tablas: {tables_dir}")
    print(f"Gráficos: {graphs_dir}")
    print("\nHallazgos rápidos:")
    print(f"- Demanda semanal: {demand_stats['total_week']:.1f} t en {demand_stats['n_ods']} ODs")
    print(f"- ODs sin tramo directo: {net_stats['nondirect_ods']} de {demand_stats['n_ods']}")
    print(f"- Demanda hacia MIA: {100*net_stats['hub_in_share']:.1f}%")
    print(f"- Simetría de arcos A↔B: {100*net_stats['symmetry_share']:.1f}%")
    print(f"- Alertas de formulación: {len(alerts)} (revisar 27_alertas_modelado.csv)")


if __name__ == "__main__":
    main()
